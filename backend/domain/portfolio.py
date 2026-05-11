from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from io import BytesIO
import json
import re
from pathlib import Path
from typing import Any

import openpyxl


ISIN_TICKERS = {
    "DE0007074007": "KWS.DE",
    "ES0144580Y14": "IBE.MC",
    "ES0173516115": "REP.MC",
    "US91324P1021": "UNH",
    "FR0000121014": "MC.PA",
    "US7427181091": "PG",
    "US9311421039": "WMT",
    "ES0157261019": "ROVI.MC",
    "XFC000A2YY6Q": "BTC-EUR",
}

FUND_NAME_MAP = {
    "INDEX MSCI EUROPE": "LU0389811539",
    "AMUNDI IX MSCI EUROPE": "LU0389811539",
    "INDEX MSCI EM": "LU0996175948",
    "AMUNDI INDEX MSCI EMERG": "LU0996175948",
    "SISF FRONTIER": "LU0968301142",
    "SCHRODER ISF FRONTIER": "LU0968301142",
    "ISHARES JAPAN": "IE00B6RVWW34",
    "ISHRS DVLP RL STTE": "IE00B83YJG36",
    "ISHARE DEVLP RL ESTATE": "IE00B83YJG36",
    "VANGUARD GLOB SMALL": "IE00B42W4L06",
    "VANGUARD US 500": "IE0032126645",
}


@dataclass(frozen=True)
class UploadedInvestmentFile:
    filename: str
    content: bytes


def build_snapshot_from_files(files: list[UploadedInvestmentFile]) -> dict[str, Any]:
    positions: list[dict[str, Any]] = []
    transactions: list[dict[str, Any]] = []
    files_summary: list[dict[str, Any]] = []

    for file in files:
        parsed = parse_investment_file(file.filename, file.content)
        positions.extend(parsed["positions"])
        transactions.extend(parsed["transactions"])
        files_summary.append(
            {
                "filename": file.filename,
                "kind": parsed["kind"],
                "positions": len(parsed["positions"]),
                "transactions": len(parsed["transactions"]),
            }
        )

    return build_snapshot(positions, transactions, files_summary)


def parse_investment_file(filename: str, content: bytes) -> dict[str, Any]:
    lowered = filename.lower()
    if lowered.endswith(".xlsx"):
        if "portfolio" in lowered:
            return parse_degiro_portfolio_xlsx(content, filename)
        if "account" in lowered:
            return parse_degiro_account_xlsx(content, filename)
        if "movimientos" in lowered or "myinvestor" in lowered:
            return parse_myinvestor_movements_xlsx(content, filename)
    if lowered.endswith(".pdf"):
        text = extract_pdf_text(content)
        if "trade republic" in text.lower() and "extracto del patrimonio neto" in text.lower():
            return parse_trade_republic_net_worth_pdf(text, filename)
        if "trade republic" in text.lower() and "transacciones de cuenta" in text.lower():
            return parse_trade_republic_account_pdf(text, filename)
        if "myinvestor" in text.lower() or "posición integrada" in text.lower() or "posicion integrada" in text.lower():
            return parse_myinvestor_statement_pdf(text, filename)
    return {"kind": "desconocido", "positions": [], "transactions": []}


def extract_pdf_text(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def parse_degiro_portfolio_xlsx(content: bytes, filename: str) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    positions = []

    for row in rows[1:]:
        product, symbol, quantity, price, value_currency, local_value, eur_value = pad(row, 7)
        if not product:
            continue
        product = str(product).strip()
        eur_value = parse_number(eur_value)
        if eur_value is None:
            continue
        asset_type = "cash" if product.upper().startswith("CASH") else ("crypto" if "BITCOIN" in product.upper() else "stock")
        positions.append(
            {
                "broker": "DeGiro",
                "source": filename,
                "asset_type": asset_type,
                "isin": str(symbol).strip() if symbol else None,
                "ticker": ISIN_TICKERS.get(str(symbol).strip()) if symbol else None,
                "name": product,
                "quantity": parse_number(quantity),
                "price": parse_number(price),
                "currency": str(value_currency).strip() if value_currency else "EUR",
                "current_value": eur_value,
            }
        )

    return {"kind": "degiro_portfolio", "positions": positions, "transactions": []}


def parse_degiro_account_xlsx(content: bytes, filename: str) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    transactions = []
    header = [str(value or "").strip().lower() for value in rows[0]]
    indexes = {name: index for index, name in enumerate(header)}

    for row in rows[1:]:
        description = str(row[indexes.get("descripción", 5)] or "").strip()
        if not description:
            continue
        isin = row[indexes.get("isin", 4)]
        isin = str(isin).strip() if isin else None
        amount = parse_number(row[8] if len(row) > 8 else None)
        currency = str(row[7] or "EUR").strip() if len(row) > 7 and row[7] else "EUR"
        date = normalize_date(row[indexes.get("fecha", 0)])
        kind = classify_description(description)
        quantity = extract_quantity(description)
        price = extract_trade_price(description)
        if kind not in {"buy", "dividend", "tax", "fee", "fx", "deposit"}:
            continue
        transactions.append(
            {
                "broker": "DeGiro",
                "source": filename,
                "date": date,
                "kind": kind,
                "isin": isin,
                "ticker": ISIN_TICKERS.get(isin or ""),
                "name": str(row[indexes.get("producto", 3)] or "").strip() or extract_name(description),
                "quantity": quantity,
                "price": price,
                "amount": amount,
                "currency": currency,
                "description": description,
            }
        )

    return {"kind": "degiro_account", "positions": [], "transactions": transactions}


def parse_myinvestor_movements_xlsx(content: bytes, filename: str) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    transactions = []

    for row in sheet.iter_rows(min_row=11, values_only=True):
        _empty, date, value_date, description, _blank, amount, balance = pad(row, 7)
        description = str(description or "").strip()
        amount = parse_number(amount)
        if date is None or amount is None:
            continue
        kind = "other"
        if "Aportacion" in description:
            kind = "deposit"
        elif "PERIODO" in description:
            kind = "interest"
        elif "COM." in description or "IVA" in description or "EFECTIVO-EUR" in description:
            kind = "fee"
        elif amount < 0 and description:
            kind = "buy"
        isin = infer_fund_isin(description)
        transactions.append(
            {
                "broker": "MyInvestor",
                "source": filename,
                "date": normalize_date(date),
                "value_date": normalize_date(value_date),
                "kind": kind,
                "isin": isin,
                "ticker": None,
                "name": clean_fund_name(description),
                "quantity": None,
                "price": None,
                "amount": amount,
                "currency": "EUR",
                "balance": parse_number(balance),
                "description": description,
            }
        )

    return {"kind": "myinvestor_movements", "positions": [], "transactions": transactions}


def parse_trade_republic_net_worth_pdf(text: str, filename: str) -> dict[str, Any]:
    positions = []
    cash_match = re.search(r"Cuenta corriente\s+([\d.,]+)\s+EUR", text)
    if cash_match:
        positions.append(
            {
                "broker": "Trade Republic",
                "source": filename,
                "asset_type": "cash",
                "isin": None,
                "ticker": None,
                "name": "Cuenta corriente",
                "quantity": None,
                "price": None,
                "currency": "EUR",
                "current_value": parse_euro(cash_match.group(1)),
            }
        )

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    isin_indexes = [(index, match.group(1)) for index, line in enumerate(lines) if (match := re.search(r"ISIN:\s*([A-Z0-9]{12})", line))]
    for index, isin in isin_indexes:
        quantity, name = extract_trade_republic_position_header(lines, index)
        price, value = extract_trade_republic_position_values(lines, index)
        if value is None:
            continue
        positions.append(
            {
                "broker": "Trade Republic",
                "source": filename,
                "asset_type": "stock",
                "isin": isin,
                "ticker": ISIN_TICKERS.get(isin),
                "name": name or isin,
                "quantity": quantity,
                "price": price,
                "currency": "EUR",
                "current_value": value,
            }
        )

    return {"kind": "trade_republic_net_worth", "positions": positions, "transactions": []}


def parse_trade_republic_account_pdf(text: str, filename: str) -> dict[str, Any]:
    transactions = []
    compact = normalize_spaces(text)
    operation_pattern = re.compile(
        r"(?P<day>\d{2}|[0-3]?\d)\s+(?P<month>[a-z]{3})(?:\s+\d{4})?\s+(?P<type>Operar|Rentabilidad|Interés|Bonificación|Transferencia)\s+"
        r"(?P<description>.*?)(?=(?:\d{2}|[0-3]?\d)\s+[a-z]{3}(?:\s+\d{4})?\s+(?:Operar|Rentabilidad|Interés|Bonificación|Transferencia|Transacción|Invitación)|RESUMEN DEL BALANCE|TRADE REPUBLIC BANK|$)",
        re.I,
    )
    for match in operation_pattern.finditer(compact):
        raw_type = match.group("type")
        description = match.group("description").strip()
        isin_match = re.search(r"([A-Z]{2}[A-Z0-9]{10})", description)
        isin = isin_match.group(1) if isin_match else None
        amount = extract_transaction_euro_amount(description)
        if amount is None:
            continue
        kind = "other"
        if raw_type.lower() == "operar" and ("buy trade" in description.lower() or "savings plan execution" in description.lower()):
            kind = "buy"
        elif raw_type.lower() == "rentabilidad" and "dividend" in description.lower():
            kind = "dividend"
        elif raw_type.lower() == "interés":
            kind = "interest"
        elif raw_type.lower() == "bonificación":
            kind = "bonus"
        elif raw_type.lower() == "transferencia":
            kind = "transfer"
        if kind == "other":
            continue
        quantity = extract_quantity(description)
        transactions.append(
            {
                "broker": "Trade Republic",
                "source": filename,
                "date": parse_spanish_date(match.group("day"), match.group("month")),
                "kind": kind,
                "isin": isin,
                "ticker": ISIN_TICKERS.get(isin or ""),
                "name": extract_trade_republic_name(description, isin),
                "quantity": quantity,
                "price": None,
                "amount": -amount if kind == "buy" else amount,
                "currency": "EUR",
                "description": description,
            }
        )

    return {"kind": "trade_republic_account", "positions": [], "transactions": transactions}


def parse_myinvestor_statement_pdf(text: str, filename: str) -> dict[str, Any]:
    positions = []
    cash_match = re.search(r"Efectivo\s+([\d.,]+)\s*€", text)
    investment_match = re.search(r"Inversión\s+([\d.,]+)\s*€", text)
    if cash_match:
        positions.append(
            {
                "broker": "MyInvestor",
                "source": filename,
                "asset_type": "cash",
                "isin": None,
                "ticker": None,
                "name": "Efectivo MyInvestor",
                "quantity": None,
                "price": None,
                "currency": "EUR",
                "current_value": parse_euro(cash_match.group(1)),
            }
        )
    position_lines = extract_myinvestor_position_lines(text)
    isin_indexes = [(index, match.group(1)) for index, line in enumerate(position_lines) if (match := re.search(r"([A-Z]{2}[A-Z0-9]{10})", line))]
    for position_number, (index, isin) in enumerate(isin_indexes):
        next_index = isin_indexes[position_number + 1][0] if position_number + 1 < len(isin_indexes) else len(position_lines)
        segment = normalize_spaces(" ".join(position_lines[index:next_index]))
        values = re.findall(r"([\d.]+,\d{4})\s*€", segment)
        if not values:
            continue
        name, quantity = extract_myinvestor_position_name_quantity(segment, isin)
        positions.append(
            {
                "broker": "MyInvestor",
                "source": filename,
                "asset_type": "fund",
                "isin": isin,
                "ticker": None,
                "name": name or isin,
                "quantity": quantity,
                "price": None,
                "currency": "EUR",
                "current_value": parse_euro(values[-1]),
            }
        )
    if investment_match and not any(position["asset_type"] == "fund" for position in positions):
        positions.append(
            {
                "broker": "MyInvestor",
                "source": filename,
                "asset_type": "fund",
                "isin": None,
                "ticker": None,
                "name": "Inversión MyInvestor",
                "quantity": None,
                "price": None,
                "currency": "EUR",
                "current_value": parse_euro(investment_match.group(1)),
            }
        )
    return {"kind": "myinvestor_statement", "positions": positions, "transactions": []}


def build_snapshot(
    raw_positions: list[dict[str, Any]],
    transactions: list[dict[str, Any]],
    files_summary: list[dict[str, Any]],
) -> dict[str, Any]:
    positions = merge_positions(raw_positions)
    cost_by_key = defaultdict(float)
    dividends_by_key, dividends_by_broker = calculate_effective_dividends(transactions)
    fees_by_broker = defaultdict(float)
    interest_by_broker = defaultdict(float)
    deposits_by_broker = defaultdict(float)

    for transaction in transactions:
        key = position_key(transaction)
        amount = transaction.get("amount") or 0.0
        if transaction["kind"] == "buy" and key:
            cost_by_key[key] += abs(amount)
        elif transaction["kind"] == "fee":
            fees_by_broker[transaction["broker"]] += abs(amount)
        elif transaction["kind"] == "interest":
            interest_by_broker[transaction["broker"]] += amount
        elif transaction["kind"] == "deposit":
            deposits_by_broker[transaction["broker"]] += amount

    enriched_positions = []
    for position in positions:
        key = position_key(position)
        cost = cost_by_key.get(key)
        dividends = dividends_by_key.get(key, 0.0)
        value = position.get("current_value") or 0.0
        pnl = value - cost if cost else None
        pnl_pct = (pnl / cost * 100.0) if cost else None
        enriched_positions.append(
            {
                **position,
                "cost": round(cost, 4) if cost else None,
                "dividends": round(dividends, 4),
                "unrealized_gain": round(pnl, 4) if pnl is not None else None,
                "unrealized_gain_pct": round(pnl_pct, 2) if pnl_pct is not None else None,
                "horizon": infer_horizon(position),
            }
        )

    summary = summarize(enriched_positions, transactions, fees_by_broker, interest_by_broker, deposits_by_broker, dividends_by_broker)
    return {
        "summary": summary,
        "positions": sorted(enriched_positions, key=lambda item: (item["broker"], item["asset_type"], item["name"])),
        "transactions": transactions,
        "brokers": summarize_brokers(enriched_positions, transactions, dividends_by_broker),
        "files": files_summary,
        "warnings": build_warnings(enriched_positions),
    }


def merge_positions(positions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    merged: dict[tuple[str, str, str], dict[str, Any]] = {}
    for position in positions:
        key = (position["broker"], position.get("isin") or position["name"], position["asset_type"])
        if key not in merged:
            merged[key] = {**position}
            continue
        current = merged[key]
        current["quantity"] = add_optional(current.get("quantity"), position.get("quantity"))
        current["current_value"] = add_optional(current.get("current_value"), position.get("current_value")) or 0.0
    return list(merged.values())


def summarize(
    positions: list[dict[str, Any]],
    transactions: list[dict[str, Any]],
    fees_by_broker: dict[str, float],
    interest_by_broker: dict[str, float],
    deposits_by_broker: dict[str, float],
    dividends_by_broker: dict[str, float],
) -> dict[str, Any]:
    total_value = sum(position.get("current_value") or 0.0 for position in positions)
    cash = sum(position.get("current_value") or 0.0 for position in positions if position["asset_type"] == "cash")
    invested = total_value - cash
    known_cost = sum(position.get("cost") or 0.0 for position in positions if position["asset_type"] != "cash")
    known_gain = sum(position.get("unrealized_gain") or 0.0 for position in positions if position.get("unrealized_gain") is not None)
    dividends = sum(dividends_by_broker.values())
    interest = sum(interest_by_broker.values())
    fees = sum(fees_by_broker.values())
    return {
        "total_value": round(total_value, 2),
        "cash": round(cash, 2),
        "invested": round(invested, 2),
        "known_cost": round(known_cost, 2),
        "known_unrealized_gain": round(known_gain, 2),
        "known_unrealized_gain_pct": round(known_gain / known_cost * 100.0, 2) if known_cost else None,
        "dividends": round(dividends, 2),
        "interest": round(interest, 2),
        "fees": round(fees, 2),
        "deposits": round(sum(deposits_by_broker.values()), 2),
        "transactions": len(transactions),
        "positions": len(positions),
    }


def summarize_brokers(
    positions: list[dict[str, Any]],
    transactions: list[dict[str, Any]],
    dividends_by_broker: dict[str, float],
) -> list[dict[str, Any]]:
    brokers = sorted({position["broker"] for position in positions} | {transaction["broker"] for transaction in transactions})
    rows = []
    for broker in brokers:
        broker_positions = [position for position in positions if position["broker"] == broker]
        broker_transactions = [transaction for transaction in transactions if transaction["broker"] == broker]
        total = sum(position.get("current_value") or 0.0 for position in broker_positions)
        cash = sum(position.get("current_value") or 0.0 for position in broker_positions if position["asset_type"] == "cash")
        cost = sum(position.get("cost") or 0.0 for position in broker_positions if position["asset_type"] != "cash")
        gain = sum(position.get("unrealized_gain") or 0.0 for position in broker_positions if position.get("unrealized_gain") is not None)
        rows.append(
            {
                "broker": broker,
                "total_value": round(total, 2),
                "cash": round(cash, 2),
                "invested": round(total - cash, 2),
                "known_cost": round(cost, 2),
                "known_unrealized_gain": round(gain, 2),
                "known_unrealized_gain_pct": round(gain / cost * 100.0, 2) if cost else None,
                "dividends": round(dividends_by_broker.get(broker, 0.0), 2),
                "interest": round(sum(t.get("amount") or 0.0 for t in broker_transactions if t["kind"] == "interest"), 2),
            }
        )
    return rows


def calculate_effective_dividends(transactions: list[dict[str, Any]]) -> tuple[dict[tuple[str, str], float], dict[str, float]]:
    by_key: defaultdict[tuple[str, str], float] = defaultdict(float)
    by_broker: defaultdict[str, float] = defaultdict(float)

    for broker in sorted({transaction["broker"] for transaction in transactions}):
        broker_transactions = [transaction for transaction in transactions if transaction["broker"] == broker]
        if broker == "Trade Republic":
            add_trade_republic_dividends(broker_transactions, by_key, by_broker)
        elif broker == "DeGiro":
            add_degiro_dividends(broker_transactions, by_key, by_broker)
        else:
            add_direct_dividends(broker_transactions, by_key, by_broker)

    return dict(by_key), dict(by_broker)


def add_direct_dividends(
    transactions: list[dict[str, Any]],
    by_key: defaultdict[tuple[str, str], float],
    by_broker: defaultdict[str, float],
) -> None:
    for transaction in transactions:
        if transaction["kind"] != "dividend":
            continue
        amount = transaction.get("amount") or 0.0
        add_effective_dividend(transaction, amount, by_key, by_broker)


def add_trade_republic_dividends(
    transactions: list[dict[str, Any]],
    by_key: defaultdict[tuple[str, str], float],
    by_broker: defaultdict[str, float],
) -> None:
    grouped: defaultdict[tuple[str | None, float], list[dict[str, Any]]] = defaultdict(list)
    for transaction in transactions:
        if transaction["kind"] == "dividend":
            grouped[(transaction.get("isin"), round(abs(transaction.get("amount") or 0.0), 2))].append(transaction)

    for (_isin, amount), dividend_rows in grouped.items():
        if not dividend_rows or amount == 0:
            continue
        net_amount = amount if len(dividend_rows) % 2 else 0.0
        if net_amount:
            add_effective_dividend(dividend_rows[-1], net_amount, by_key, by_broker)


def add_degiro_dividends(
    transactions: list[dict[str, Any]],
    by_key: defaultdict[tuple[str, str], float],
    by_broker: defaultdict[str, float],
) -> None:
    assigned_dividends: set[int] = set()
    for index, transaction in enumerate(transactions):
        if transaction["kind"] != "fx" or transaction.get("currency") != "EUR" or (transaction.get("amount") or 0.0) <= 0:
            continue
        candidate = find_nearby_foreign_dividend(transactions, index, assigned_dividends)
        if candidate is None:
            continue
        assigned_dividends.add(candidate)
        add_effective_dividend(transactions[candidate], transaction.get("amount") or 0.0, by_key, by_broker)

    for index, transaction in enumerate(transactions):
        if index in assigned_dividends or transaction["kind"] != "dividend":
            continue
        if transaction.get("currency") == "EUR":
            add_effective_dividend(transaction, transaction.get("amount") or 0.0, by_key, by_broker)


def find_nearby_foreign_dividend(
    transactions: list[dict[str, Any]],
    fx_index: int,
    assigned_dividends: set[int],
) -> int | None:
    candidates = []
    for index, transaction in enumerate(transactions):
        if index in assigned_dividends:
            continue
        if transaction["kind"] != "dividend" or transaction.get("currency") == "EUR" or (transaction.get("amount") or 0.0) <= 0:
            continue
        distance = abs(index - fx_index)
        if distance <= 6:
            candidates.append((distance, index))
    return min(candidates)[1] if candidates else None


def add_effective_dividend(
    transaction: dict[str, Any],
    amount: float,
    by_key: defaultdict[tuple[str, str], float],
    by_broker: defaultdict[str, float],
) -> None:
    if not amount:
        return
    key = position_key(transaction)
    if key:
        by_key[key] += amount
    by_broker[transaction["broker"]] += amount


def build_warnings(positions: list[dict[str, Any]]) -> list[str]:
    unknown_cost = [position["name"] for position in positions if position["asset_type"] != "cash" and not position.get("cost")]
    if not unknown_cost:
        return []
    return [f"Hay {len(unknown_cost)} posiciones sin coste histórico detectado; su rentabilidad queda pendiente."]


def position_key(item: dict[str, Any]) -> tuple[str, str] | None:
    identifier = item.get("isin") or item.get("name")
    if not identifier:
        return None
    return (item["broker"], identifier)


def infer_horizon(position: dict[str, Any]) -> str:
    if position["asset_type"] == "cash":
        return "corto"
    if position["asset_type"] in {"fund", "stock"}:
        return "largo"
    return "medio"


def parse_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("€", "").replace("EUR", "").replace("USD", "").strip()
    if not text:
        return None
    text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def parse_euro(value: str) -> float | None:
    return parse_number(value)


def normalize_date(value: Any) -> str | None:
    if value is None:
        return None
    return str(value).strip()


def parse_spanish_date(day: str, month: str) -> str:
    months = {
        "ene": "01",
        "feb": "02",
        "mar": "03",
        "abr": "04",
        "may": "05",
        "jun": "06",
        "jul": "07",
        "ago": "08",
        "sep": "09",
        "oct": "10",
        "nov": "11",
        "dic": "12",
    }
    return f"2026-{months.get(month.lower(), '01')}-{int(day):02d}"


def pad(row: tuple[Any, ...], length: int) -> tuple[Any, ...]:
    return tuple(row) + (None,) * max(0, length - len(row))


def normalize_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def classify_description(description: str) -> str:
    lowered = description.lower()
    if lowered.startswith("compra"):
        return "buy"
    if "retención" in lowered or "tax" in lowered:
        return "tax"
    if "dividendo" in lowered:
        return "dividend"
    if "costes" in lowered or "comisión" in lowered:
        return "fee"
    if "cambio de divisa" in lowered:
        return "fx"
    if "deposit" in lowered:
        return "deposit"
    return "other"


def extract_quantity(description: str) -> float | None:
    quantity_match = re.search(r"quantity:\s*([\d.,]+)", description, re.I)
    if quantity_match:
        return parse_decimal_quantity(quantity_match.group(1))
    purchase_match = re.search(r"Compra\s+([\d.,]+)", description, re.I)
    if purchase_match:
        return parse_number(purchase_match.group(1))
    return None


def extract_trade_price(description: str) -> float | None:
    match = re.search(r"@([\d.,]+)", description)
    if not match:
        return None
    return parse_number(match.group(1))


def extract_trade_republic_position_header(lines: list[str], isin_index: int) -> tuple[float | None, str | None]:
    search_start = max(0, isin_index - 8)
    quantity_index = None
    quantity = None
    for index in range(isin_index - 1, search_start - 1, -1):
        line = lines[index]
        if "unidades" in line.lower() or re.fullmatch(r"\d+(?:,\d+)?", line):
            quantity_index = index
            quantity_match = re.search(r"\d+(?:,\d+)?", line)
            quantity = parse_number(quantity_match.group(0)) if quantity_match else None
            break

    if quantity_index is None:
        return None, None

    first_line = re.sub(r"^\d+(?:,\d+)?\s*(?:unidades)?\s*", "", lines[quantity_index], flags=re.I).strip()
    name_parts = [first_line] if first_line else []
    for line in lines[quantity_index + 1 : isin_index]:
        if line.lower() == "unidades":
            continue
        name_parts.append(line)
    return quantity, normalize_spaces(" ".join(name_parts)) or None


def extract_trade_republic_position_values(lines: list[str], isin_index: int) -> tuple[float | None, float | None]:
    price = None
    for index in range(isin_index + 1, min(len(lines), isin_index + 10)):
        if re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", lines[index]):
            previous_numbers = [parse_euro(line) for line in lines[isin_index + 1 : index] if parse_euro(line) is not None]
            price = previous_numbers[-1] if previous_numbers else None
            for value_line in lines[index + 1 : min(len(lines), index + 4)]:
                value = parse_euro(value_line)
                if value is not None:
                    return price, value
    return price, None


def extract_myinvestor_position_lines(text: str) -> list[str]:
    if "Posiciones" not in text:
        return []
    section = text.split("Posiciones", 1)[1].split("Tarjetas", 1)[0]
    return [line.strip() for line in section.splitlines() if line.strip()]


def extract_myinvestor_position_name_quantity(segment: str, isin: str) -> tuple[str | None, float | None]:
    content = segment.split(isin, 1)[1].strip() if isin in segment else segment
    match = re.search(r"(?P<name>.+?)\s+EUR\s+(?P<quantity>\d+(?:[.,]\d+)?)\s+[\d.]+,\d{4}\s*€", content)
    if not match:
        return normalize_spaces(content), None
    return normalize_spaces(match.group("name")), parse_decimal_quantity(match.group("quantity"))


def parse_decimal_quantity(value: str) -> float | None:
    try:
        return float(value.replace(",", "."))
    except ValueError:
        return None


def extract_name(description: str) -> str:
    match = re.search(r"Compra\s+[\d.,]+\s+(.+?)@", description)
    return normalize_spaces(match.group(1)) if match else description


def infer_fund_isin(description: str) -> str | None:
    upper = description.upper()
    for fragment, isin in FUND_NAME_MAP.items():
        if fragment in upper:
            return isin
    return None


def clean_fund_name(description: str) -> str:
    return normalize_spaces(description.split("@", 1)[0])


def extract_last_euro_amount(description: str) -> float | None:
    amounts = re.findall(r"([\d.]+,\d{2})\s*€", description)
    if not amounts:
        return None
    return parse_euro(amounts[-1])


def extract_transaction_euro_amount(description: str) -> float | None:
    amounts = re.findall(r"([\d.]+,\d{2})\s*€", description)
    if not amounts:
        return None
    return parse_euro(amounts[-2] if len(amounts) > 1 else amounts[0])


def extract_trade_republic_name(description: str, isin: str | None) -> str:
    if not isin:
        return description[:80]
    after = description.split(isin, 1)[-1]
    after = re.split(r",\s*quantity:|quantity:", after, maxsplit=1, flags=re.I)[0]
    after = after.replace("Buy trade", "").replace("Savings plan execution", "")
    return normalize_spaces(after) or isin


def add_optional(left: float | None, right: float | None) -> float | None:
    if left is None:
        return right
    if right is None:
        return left
    return left + right


def save_snapshot(path: Path, snapshot: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")


def load_snapshot(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))
